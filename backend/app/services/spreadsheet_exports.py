"""Native Excel reports, built in memory with openpyxl. Never persists inputs."""
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.data_source import AxDataSource, StrRef, StrData, StrVal
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from app.schemas_exports import HealthExport, HistoryExport

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
INK, ACCENT, BLUE, MUTED = "203238", "087F75", "3976C6", "64748B"
NUMBER, DECIMAL = '#,##0;[Red](#,##0);0', '#,##0.0;[Red](#,##0.0);0'
HEALTH = {
    "stockout_risk": ("Stockout risk", "Check incoming delivery dates and current demand."),
    "reorder_review": ("Review a reorder", "Check open orders, committed stock and supplier minimums."),
    "excess_stock": ("Above target cover", "Review demand changes and planned promotions before buying more."),
    "no_recent_sales": ("No recorded sales", "Check whether this is new, seasonal or previously unavailable."),
    "within_range": ("Within chosen range", "Keep monitoring sales and lead times."),
}


def literal(cell, value):
    """Identifiers remain text, including leading zeros and formula-like names."""
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def currency_format(code):
    return f'"{code} "#,##0.00;[Red]("{code} "#,##0.00);"{code} "0.00'


def _sheet(wb, name, title, context):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_properties.tabColor = ACCENT
    ws.sheet_format.defaultRowHeight = 21
    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 23
    ws.column_dimensions["D"].width = 3
    literal(ws["B2"], title).font = Font(name="Arial", size=16, bold=True, color=INK)
    ws.row_dimensions[2].height = 28
    for col in range(2, 15):
        ws.cell(3, col).border = Border(bottom=Side(style="thin", color=ACCENT))
    literal(ws["B4"], context).font = Font(name="Arial", size=10, color=MUTED)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.oddFooter.center.text = "Skubase | Page &P of &N"
    return ws


def _workbook(title, context):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Skubase"
    wb.properties.title = title
    wb.calculation.fullCalcOnLoad = True
    summary = _sheet(wb, "Summary", title, context)
    summary.sheet_properties.tabColor = INK
    return wb, summary


def _table(wb, name, title, columns, rows, context=""):
    ws = _sheet(wb, name, title, context)
    for j, (label, width, fmt) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
        c = literal(ws.cell(5, j), label)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 34
    for i, row in enumerate(rows, 6):
        for j, value in enumerate(row, 1):
            c = literal(ws.cell(i, j), value)
            c.font = Font(name="Arial", size=10, color=INK)
            c.number_format = columns[j - 1][2]
            c.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left", wrap_text=isinstance(value, str))
        # Full labels remain readable without squeezing the font.
        longest = max((len(str(v)) / max(columns[j][1] - 2, 1) for j, v in enumerate(row) if v is not None), default=1)
        ws.row_dimensions[i].height = min(150, max(24, 15 * (int(longest) + 1)))
    end = max(5, len(rows) + 5)
    ref = f"A5:{get_column_letter(len(columns))}{end}"
    if rows:
        table = Table(displayName=name.replace(" ", "") + "Data", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.auto_filter.ref = ref
    ws.freeze_panes = "B6"
    ws.print_title_rows = "1:5"
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(len(columns))}{max(end, 7)}"
    if not rows:
        literal(ws["A6"], "No records to export.")
    return ws


def _metric(ws, row, label, formula, fmt=NUMBER):
    literal(ws.cell(row, 2), label).font = Font(name="Arial", size=10, color=INK)
    c = ws.cell(row, 3, formula)
    c.font = Font(name="Arial", size=11, bold=True, color=ACCENT)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    c.fill = PatternFill("solid", fgColor="EAF5F3")
    ws.row_dimensions[row].height = 26


def _note(ws, row, text):
    literal(ws.cell(row, 2), text).font = Font(name="Arial", size=10, color=MUTED)


def _bar(ws, start, end, title, anchor, unit="SKUs", population=5, labels=None):
    if end < start:
        return
    chart = BarChart()
    chart.type, chart.style, chart.title = "bar", 13, title
    chart.add_data(Reference(ws, min_col=3, min_row=start, max_row=end))
    labels = labels if labels is not None else [str(ws.cell(row, 2).value or "Unknown") for row in range(start, end + 1)]
    labels = [ILLEGAL_CHARACTERS_RE.sub("", label) for label in labels]
    cache = StrData(ptCount=len(labels), pt=[StrVal(idx=i, v=label) for i, label in enumerate(labels)])
    chart.series[0].cat = AxDataSource(strRef=StrRef(f=f"'{ws.title}'!$B${start}:$B${end}", strCache=cache))
    chart.legend = None
    # openpyxl keeps x_axis categorical and y_axis numeric even for horizontal bars.
    chart.y_axis.title = unit
    chart.y_axis.numFmt = '#,##0'
    chart.y_axis.scaling.min = 0
    if unit in {"SKUs", "Leads"}:
        chart.y_axis.majorUnit = max(1, population // 5)
    chart.x_axis.scaling.orientation = "maxMin"
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.crosses = "max"
    chart.layout = Layout(manualLayout=ManualLayout(x=.3, y=.15, w=.65, h=.7, xMode="factor", yMode="factor"))
    chart.width, chart.height = 22, max(7.5, (end - start + 1) * .75)
    chart.series[0].graphicalProperties.solidFill = ACCENT
    chart.series[0].graphicalProperties.line.noFill = True
    ws.add_chart(chart, anchor)


def _save(wb):
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def history_workbook(payload: HistoryExport):
    points = sorted(payload.points, key=lambda p: p.date)
    context = f"{points[0].date} to {points[-1].date} | Daily snapshots (UTC) | USD"
    if payload.sample:
        context = "SAMPLE DATA | " + context
    wb, ws = _workbook("Inventory value over time", context)
    detail = _table(wb, "Daily history", "Daily inventory snapshots", [
        ("Date (UTC)", 18, "mmm d, yyyy"), ("Recorded-cost subtotal (USD)", 29, currency_format("USD")),
        ("Retail value (USD)", 26, currency_format("USD")), ("Units on hand", 22, NUMBER),
    ], [[p.date, p.cost_value, p.retail_value, p.total_units] for p in points], "Blank values are unknown. Cost is a recorded-cost subtotal, not complete capital.")
    end = len(points) + 5
    _metric(ws, 6, "Daily snapshots", f"=COUNTA('Daily history'!A6:A{end})")
    for row, label, col, fmt in [(7, "Latest recorded-cost subtotal", "B", currency_format("USD")),
                                 (8, "Latest retail value", "C", currency_format("USD")),
                                 (9, "Latest units on hand", "D", NUMBER)]:
        _metric(ws, row, label, f'=IF(ISNUMBER(\'Daily history\'!{col}{end}),\'Daily history\'!{col}{end},"Unknown")', fmt)
    _note(ws, 12, "Cost excludes unrecorded unit costs.")
    _note(ws, 13, "Retail value is not realized revenue.")
    _note(ws, 14, "Missing dates are not estimated.")
    for columns, title, anchor, fmt in [([2, 3], "Inventory value (USD)", "E6", '#,##0'), ([4], "Units on hand", "E24", NUMBER)]:
        chart = LineChart()
        chart.title, chart.style = title, 13
        chart.width, chart.height = 22, 9
        chart.display_blanks = "gap"
        chart.x_axis = DateAxis(crossAx=100)
        chart.x_axis.number_format = "mmm d"
        chart.x_axis.majorTimeUnit = "days"
        chart.x_axis.majorUnit = max(1, (len(points) + 6) // 7)
        chart.y_axis.numFmt = fmt
        chart.legend.position = "b"
        for j, col in enumerate(columns):
            chart.add_data(Reference(detail, min_col=col, min_row=5, max_row=end), titles_from_data=True)
            chart.series[-1].tx = SeriesLabel(v={2: "Recorded-cost subtotal", 3: "Retail value", 4: "Units on hand"}[col])
            chart.series[-1].graphicalProperties.line.solidFill = [ACCENT, BLUE][j]
            chart.series[-1].graphicalProperties.line.width = 25000
            if len(points) == 1:
                chart.series[-1].marker.symbol = "circle"
        chart.set_categories(Reference(detail, min_col=1, min_row=6, max_row=end))
        ws.add_chart(chart, anchor)
    ws.print_area = "B1:O43"
    return _save(wb)


def health_workbook(payload: HealthExport):
    settings, rows = payload.settings, payload.rows
    context = f"{settings.salesDays}-day sales period | {settings.targetCoverDays}-day stock target | Costs in {payload.currency}"
    if payload.sample:
        context = "SAMPLE DATA | " + context
    wb, ws = _workbook("Inventory health check", context)
    money = currency_format(payload.currency)
    detail = _table(wb, "SKU detail", "Inventory review priorities", [
        ("SKU", 26, "@"), ("Review priority", 25, "@"), ("On hand", 16, NUMBER),
        ("Units sold", 16, NUMBER), ("Days of cover", 18, DECIMAL), ("Reorder trigger (units)", 22, NUMBER),
        ("Above target (units)", 22, NUMBER), (f"Above-target cost ({payload.currency})", 26, money),
        ("Review action", 62, "@"), ("Cost currency", 16, "@"), ("Sales period (days)", 20, NUMBER),
        ("Target cover (days)", 20, NUMBER), ("Lead time (days)", 19, NUMBER), ("Safety stock (units)", 20, NUMBER),
        (f"Unit cost ({payload.currency})", 22, money),
    ], [[r.sku, HEALTH[r.status][0], r.onHand, r.unitsSold, r.daysCover,
         r.reorderPoint if r.unitsSold > 0 else None, r.excessUnits, r.excessCost,
         HEALTH[r.status][1], payload.currency, settings.salesDays, settings.targetCoverDays,
         r.leadDays, r.safetyStock, r.unitCost] for r in rows],
        "Blank values are unknown or unavailable. No recorded sales alone does not establish dead stock.")
    end = len(rows) + 5
    for i, (_, (label, _)) in enumerate(HEALTH.items(), 6):
        _metric(ws, i, label, f'=COUNTIF(\'SKU detail\'!B6:B{end},B{i})')
    _bar(ws, 6, 10, "Inventory review priorities", "E6", population=len(rows))
    _metric(ws, 14, "Total SKUs", f"=COUNTA('SKU detail'!A6:A{end})")
    _metric(ws, 15, "Known above-target cost", f'=IF(AND(C16>0,SUM(\'SKU detail\'!H6:H{end})=0),"Unknown",SUM(\'SKU detail\'!H6:H{end}))', money)
    _metric(ws, 16, "Above-target SKUs missing cost", f'=COUNTIFS(\'SKU detail\'!G6:G{end},">0",\'SKU detail\'!O6:O{end},"")')
    _note(ws, 18, "Known cost is a subtotal, not guaranteed recovery.")
    _note(ws, 19, "See SKU detail for assumptions and next actions.")
    ranked = sorted(((i + 6, r) for i, r in enumerate(rows) if r.excessCost is not None and r.excessCost > 0), key=lambda pair: pair[1].excessCost, reverse=True)[:10]
    _note(ws, 24, "Largest known above-target positions at export")
    for row, (source_row, _) in enumerate(ranked, 26):
        label = ws.cell(row, 2, f"='SKU detail'!A{source_row}")
        label.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 45
        ws.cell(row, 3, f"='SKU detail'!H{source_row}").number_format = money
    _bar(ws, 26, 25 + len(ranked), f"Above-target cost ({payload.currency})", "E24", payload.currency, labels=[r.sku for _, r in ranked])
    if not ranked:
        _note(ws, 26, "No known above-target cost to chart.")
    for label, color in [("Stockout risk", "FEE2E2"), ("Review a reorder", "FFF3CF"), ("Above target cover", "FFF3CF")]:
        detail.conditional_formatting.add(f"B6:B{end}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=PatternFill("solid", fgColor=color)))
    ws.print_area = "B1:O43"
    return _save(wb)


def leads_workbook(leads):
    generated = datetime.now(timezone.utc).strftime("%b %d, %Y %H:%M UTC")
    wb, ws = _workbook("Inventory snapshot leads", f"Generated {generated} | Captured leads, not verified customers")
    columns = [("ID", 10, NUMBER), ("Captured at (UTC)", 23, "mmm d, yyyy hh:mm"),
               ("First name", 20, "@"), ("Email", 34, "@"), ("Company", 28, "@"), ("Store URL", 38, "@"),
               ("Approximate SKU count", 23, "@"), ("Biggest issue", 28, "@"), ("Source", 22, "@"),
               ("UTM source", 22, "@"), ("UTM medium", 22, "@"), ("UTM campaign", 28, "@"),
               ("UTM content", 25, "@"), ("UTM term", 22, "@"), ("Status", 22, "@"), ("Attributed source", 25, "@")]
    values = []
    for lead in leads:
        captured = lead.created_at
        if captured and captured.tzinfo:
            captured = captured.astimezone(timezone.utc).replace(tzinfo=None)
        values.append([lead.id, captured, lead.first_name, lead.email, lead.company_name, lead.store_url,
                       lead.approximate_sku_count, lead.biggest_inventory_issue, lead.source,
                       lead.utm_source, lead.utm_medium, lead.utm_campaign, lead.utm_content, lead.utm_term, lead.status,
                       lead.utm_source or lead.source or "direct"])
    _table(wb, "Leads", "Captured lead records", columns, values, "Captured timestamps are UTC. SKU ranges remain text.")
    end = max(6, 5 + len(leads))
    statuses = ["New", "Reviewing", "Snapshot sent", "Demo booked", "Not qualified"]
    for row, status in enumerate(statuses, 6):
        _metric(ws, row, status, f'=COUNTIF(Leads!O6:O{end},B{row})')
    _bar(ws, 6, 10, "Lead status", "E6", "Leads", len(leads))
    issues = Counter(lead.biggest_inventory_issue for lead in leads)
    for row, (label, _) in enumerate(issues.most_common(), 26):
        _metric(ws, row, label, f'=SUMPRODUCT(--(Leads!H6:H{end}=B{row}))')
    _bar(ws, 26, 25 + len(issues), "Inventory problems mentioned", "E24", "Leads", len(leads))
    sources = Counter(lead.utm_source or lead.source or "direct" for lead in leads)
    for row, (label, _) in enumerate(sources.most_common(8), 46):
        _metric(ws, row, label, f'=SUMPRODUCT(--(Leads!P6:P{end}=B{row}))')
    _bar(ws, 46, 45 + min(8, len(sources)), "Acquisition sources (top 8)", "E43", "Leads", len(leads))
    _metric(ws, 14, "Captured leads", f"=COUNT(Leads!A6:A{end})")
    if not leads:
        _note(ws, 18, "No leads captured yet.")
    ws.print_area = "B1:O62"
    return _save(wb)
