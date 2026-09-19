"""Workbook contents and public formatter boundaries, using synthetic data only."""
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.api.routes import spreadsheet_exports as route
from app.schemas_exports import HealthExport, HistoryExport
from app.services.spreadsheet_exports import health_workbook, history_workbook, leads_workbook, MIME


def health_payload():
    return dict(kind="inventory_health", sample=True, currency="USD",
                settings=dict(salesDays=30, defaultLeadDays=14, targetCoverDays=60), rows=[
                    dict(sku="001", onHand=200, unitsSold=30, unitCost=12.5, leadDays=14,
                         safetyStock=0, dailySales=1, daysCover=200, reorderPoint=14,
                         excessUnits=140, excessCost=1750, status="excess_stock"),
                    dict(sku="=HYPERLINK(\"https://example.com\")", onHand=20, unitsSold=0,
                         unitCost=None, leadDays=14, safetyStock=0, dailySales=0,
                         daysCover=None, reorderPoint=0, excessUnits=None, excessCost=None,
                         status="no_recent_sales"),
                ])


def history_payload():
    return dict(kind="inventory_history", sample=True, points=[
        dict(date="2026-09-18", cost_value=None, retail_value=0, total_units=-2),
        dict(date="2026-09-17", cost_value=1234.5, retail_value=2500, total_units=40),
    ])


class SpreadsheetExportTests(unittest.TestCase):
    def test_history_retains_dates_missing_values_and_negative_balances(self):
        wb = load_workbook(BytesIO(history_workbook(HistoryExport(**history_payload()))))
        detail = wb["Daily history"]
        self.assertEqual(detail["A6"].value, datetime(2026, 9, 17))
        self.assertIsNone(detail["B7"].value)
        self.assertEqual(detail["C7"].value, 0)
        self.assertEqual(detail["D7"].value, -2)
        self.assertEqual(detail.freeze_panes, "B6")
        self.assertEqual(detail.tables["DailyhistoryData"].ref, "A5:D7")
        self.assertIn('"Unknown"', wb["Summary"]["C7"].value)
        self.assertIn("SAMPLE DATA", wb["Summary"]["B4"].value)
        charts = wb["Summary"]._charts
        self.assertEqual(len(charts), 2)
        self.assertEqual(charts[0].series[0].val.numRef.f, "'Daily history'!$B$6:$B$7")
        self.assertEqual(charts[0].series[0].cat.numRef.f, "'Daily history'!$A$6:$A$7")

    def test_health_preserves_identifiers_unknowns_and_native_charts(self):
        wb = load_workbook(BytesIO(health_workbook(HealthExport(**health_payload()))))
        detail = wb["SKU detail"]
        self.assertEqual(detail["A6"].value, "001")
        self.assertEqual(detail["A7"].data_type, "s")
        self.assertTrue(detail["A7"].value.startswith("=HYPERLINK"))
        for address in ("E7", "F7", "G7", "H7", "O7"):
            self.assertIsNone(detail[address].value, address)
        self.assertEqual(detail["H6"].value, 1750)
        self.assertIn("USD", detail["H6"].number_format)
        self.assertEqual(len(wb["Summary"]._charts), 2)
        cost_chart = wb["Summary"]._charts[1]
        self.assertEqual(cost_chart.y_axis.scaling.min, 0)
        self.assertEqual(cost_chart.series[0].cat.strRef.strCache.pt[0].v, "001")
        self.assertEqual(wb["Summary"]["C26"].value, "='SKU detail'!H6")
        self.assertIn("C16>0", wb["Summary"]["C15"].value)

    def test_zero_cost_health_does_not_make_an_empty_cost_chart(self):
        data = health_payload()
        data["rows"] = data["rows"][1:]
        wb = load_workbook(BytesIO(health_workbook(HealthExport(**data))))
        self.assertEqual(len(wb["Summary"]._charts), 1)

    def test_leads_keep_utc_dates_ranges_and_literal_contact_fields(self):
        lead = SimpleNamespace(id=1, created_at=datetime(2026, 9, 18, 14, tzinfo=timezone.utc),
            first_name="=1+1", email="test@example.com", company_name="Demo store", store_url="https://example.com",
            approximate_sku_count="75-100", biggest_inventory_issue="Reordering", source="direct",
            utm_source="*", utm_medium=None, utm_campaign=None, utm_content=None, utm_term=None, status="New")
        wb = load_workbook(BytesIO(leads_workbook([lead])))
        self.assertEqual(wb["Leads"]["B6"].value, datetime(2026, 9, 18, 14))
        self.assertEqual(wb["Leads"]["C6"].data_type, "s")
        self.assertEqual(wb["Leads"]["G6"].value, "75-100")
        self.assertEqual(len(wb["Summary"]._charts), 3)
        self.assertIn("SUMPRODUCT", wb["Summary"]["C46"].value)  # '*' is literal, not a wildcard.
        empty = load_workbook(BytesIO(leads_workbook([])))
        self.assertEqual(empty["Leads"]["A6"].value, "No records to export.")

    def test_bounds_duplicates_and_nonfinite_values(self):
        for change in (dict(points=[]), dict(points=history_payload()["points"] * 2),
                       dict(points=[dict(date="2026-09-18", cost_value=float("nan"), retail_value=0, total_units=1)])):
            with self.assertRaises(ValidationError):
                HistoryExport(**{**history_payload(), **change})
        data = health_payload()
        data["rows"] *= 501
        with self.assertRaises(ValidationError):
            HealthExport(**data)

    def test_stateless_http_download_and_fail_closed_limits(self):
        app = FastAPI()
        app.include_router(route.router)
        with TestClient(app) as client:
            response = client.post("/exports/workbook.xlsx", json=health_payload())
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["content-type"], MIME)
            self.assertIn("sample.xlsx", response.headers["content-disposition"])
            self.assertIn("no-store", response.headers["cache-control"])
            self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames, ["Summary", "SKU detail"])
            self.assertEqual(client.post("/exports/workbook.xlsx", content=b"x" * (route.MAX_BYTES + 1)).status_code, 413)
            invalid = client.post("/exports/workbook.xlsx", json={"secret_sku": "must-not-echo"})
            self.assertEqual(invalid.status_code, 422)
            self.assertNotIn("must-not-echo", invalid.text)
            with patch.object(route._slots, "acquire", return_value=False):
                self.assertEqual(client.post("/exports/workbook.xlsx", json=history_payload()).status_code, 503)
            with patch.object(route, "history_workbook", side_effect=RuntimeError("test builder failure")):
                with self.assertRaises(RuntimeError):
                    client.post("/exports/workbook.xlsx", json=history_payload())
            # A failure must release its concurrency permit.
            self.assertTrue(route._slots.acquire(blocking=False))
            self.assertTrue(route._slots.acquire(blocking=False))
            route._slots.release()
            route._slots.release()

    def test_lead_export_retains_administrator_authorization(self):
        from app.api.routes.inventory_risk_snapshot import router as leads_router
        from app.db.session import get_db_session
        app = FastAPI()
        app.include_router(leads_router)
        app.dependency_overrides[get_db_session] = lambda: None
        with patch.dict("os.environ", {"ADMIN_BOOTSTRAP_TOKEN": "test-only-token"}), TestClient(app) as client:
            self.assertEqual(client.get("/inventory-risk-snapshot/leads/export.xlsx").status_code, 403)


if __name__ == "__main__":
    unittest.main()
